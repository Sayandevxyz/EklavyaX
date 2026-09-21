from __future__ import annotations

import re
import secrets
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile

from app.core.security import require_role
from app.db import models
from app.services.ai_service import get_explanation

router = APIRouter(prefix="/api/study-materials", tags=["Smart Study Material"])
UPLOAD_DIR = Path(__file__).resolve().parents[4] / "frontend" / "assets" / "uploads" / "study-materials"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md"}


def local_insights(notes: str, chapter: str) -> dict[str, Any]:
    words = re.findall(r"[A-Za-z][A-Za-z-]{3,}", notes.lower())
    stopwords = {"this", "that", "with", "from", "into", "will", "have", "about", "chapter", "using"}
    concepts = []
    for word in words:
        if word not in stopwords and word not in concepts:
            concepts.append(word)
        if len(concepts) == 6:
            break
    formulas = re.findall(r"[^.!?\n]*(?:=|→|∫|√|²)[^.!?\n]*", notes)
    return {
        "summary": f"Study {chapter or 'this material'} in focused passes: definitions first, then worked examples, then recall practice.",
        "important_concepts": concepts or [chapter or "Core definitions", "Worked examples", "Key vocabulary"],
        "formula_sheet": [formula.strip() for formula in formulas[:8]],
        "chapter_notes": [line.strip() for line in notes.splitlines() if line.strip()][:8],
    }


def extract_file_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if extension in {".ppt", ".pptx"}:
        from pptx import Presentation

        presentation = Presentation(BytesIO(content))
        lines = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
        return "\n".join(lines)
    return content.decode("utf-8", errors="ignore")


@router.post("/analyze")
async def analyze_material(
    file: UploadFile | None = File(None),
    subject: str = Form("General"),
    chapter: str = Form(""),
    notes: str = Form(""),
    current_user: models.User = Depends(require_role("student")),
):
    if not file and not notes.strip():
        raise HTTPException(status_code=400, detail="Upload a PDF/PPT or paste notes first.")

    filename = None
    extracted_notes = notes.strip()
    if file:
        extension = Path(file.filename or "").suffix.lower()
        if extension not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=400, detail="Supported files: PDF, PPTX, TXT, or Markdown notes. Save legacy PPT files as PPTX first.")
        content = await file.read()
        if len(content) > 20 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Study material must be smaller than 20MB.")
        filename = f"material_{secrets.token_hex(6)}{extension}"
        (UPLOAD_DIR / filename).write_bytes(content)
        try:
            extracted_notes = extract_file_text(file.filename or filename, content).strip()
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Could not read this {extension[1:].upper()} file: {exc}") from exc

    insights = local_insights(extracted_notes, chapter)
    if extracted_notes:
        try:
            insights["summary"] = await get_explanation(
                f"Create a concise study summary for {subject}, chapter {chapter or 'unspecified'} from this uploaded study material. Extract important concepts, formulas, and chapter-wise notes.\n{extracted_notes[:12000]}",
                "Simple English",
            )
        except Exception:
            pass

    return {
        "status": "analyzed",
        "filename": file.filename if file else None,
        "stored_file": f"/assets/uploads/study-materials/{filename}" if filename else None,
        "subject": subject,
        "chapter": chapter,
        **insights,
    }
